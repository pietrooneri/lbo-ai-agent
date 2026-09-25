import json

import anthropic
import httpx2
import openpyxl
import pytest
from anthropic import DefaultHttpxClient
from anthropic.lib.tools import ToolError

from lbo_agent import SYSTEM_PROMPT, DealSession, LBOAgent
from lbo_engine import run_model
from test_assumption_generator import proposal


def fake_llm(description):
    return proposal(provided={"revenue_at_entry"})


@pytest.fixture
def session(tmp_path):
    s = DealSession(output_dir=str(tmp_path / "out"), llm=fake_llm)
    s.generate_base_case("Italian industrial, EUR 400m revenue")
    return s


# --- deterministic tools -------------------------------------------------------

def test_base_case_results_come_from_engine(session):
    base = session.scenarios["base"].assumptions
    assert session.compare_scenarios()["scenarios"][0]["irr"] == round(run_model(base)["returns"]["irr"], 4)


def test_post_close_downside_keeps_price_and_debt(session):
    base = session.scenarios["base"].assumptions
    res = session.run_scenario("downside", {"ebitda_margin": 0.12, "revenue_growth": 0.01}, "margin squeeze")
    a = session.scenarios["downside"].assumptions
    assert a.entry_ebitda == base.entry_ebitda == pytest.approx(60.0)
    assert (res["entry_ev"], res["entry_debt"]) == (round(60.0 * 8.5, 1), round(60.0 * 4.75, 1))
    assert res["projected_margin"] == 0.12 and res["ltm_margin"] == 0.15
    assert res["ebitda_by_year"] == [round(y.ebitda, 1) for y in run_model(a)["years"]]
    assert res["entry_terms"].startswith("as signed")
    assert res["irr"] == round(run_model(a)["returns"]["irr"], 4)
    assert res["overrides_vs_base"] == {"ebitda_margin": 0.12, "revenue_growth": 0.01}


def test_reprice_entry_resizes_price_and_debt(session):
    res = session.run_scenario("weaker_today", {"ebitda_margin": 0.12}, "margin already lower", reprice_entry=True)
    a = session.scenarios["weaker_today"].assumptions
    assert a.entry_ebitda == pytest.approx(400 * 0.12)
    assert res["entry_ev"] == round(48.0 * 8.5, 1)
    assert res["entry_terms"].startswith("re-priced")
    assert "entry_ebitda" in res["overrides_vs_base"]


def test_revenue_override_keeps_ltm_margin(session):
    session.run_scenario("bigger", {"revenue_at_entry": 500.0}, "larger target")
    assert session.scenarios["bigger"].assumptions.entry_ebitda == pytest.approx(500 * 0.15)


def test_solving_margin_is_post_close(session):
    res = session.solve_for_target("ebitda_margin", "moic", 1.0)
    assert res["check"]["entry_ev"] == round(60.0 * 8.5, 1)    # price unchanged while margin moves


def test_changing_total_leverage_keeps_mix_and_records_it(session):
    base = session.scenarios["base"].assumptions
    res = session.run_scenario("lev", {"total_leverage_x": 3.8}, "less debt")
    a = session.scenarios["lev"].assumptions
    assert a.senior_leverage_x / a.total_leverage_x == pytest.approx(base.senior_leverage_x / base.total_leverage_x)
    assert set(res["overrides_vs_base"]) == {"total_leverage_x", "senior_leverage_x"}


def test_scenarios_can_chain(session):
    session.run_scenario("downside", {"revenue_growth": 0.0}, "flat sales")
    session.run_scenario("downside_low_exit", {"exit_ev_multiple": 7.0}, "and derating", based_on="downside")
    a = session.scenarios["downside_low_exit"].assumptions
    assert (a.revenue_growth, a.exit_ev_multiple) == (0.0, 7.0)


@pytest.mark.parametrize("overrides, message", [
    ({"growth": 0.05}, "Unknown assumption"),
    ({"entry_ebitda": 500}, "below revenue_at_entry"),
    ({"revenue_growth": 4.0}, "decimals"),
    ({"senior_leverage_x": 9.0}, "senior_leverage_x"),
])
def test_invalid_overrides_return_actionable_errors(session, overrides, message):
    with pytest.raises(ToolError) as exc:
        session.run_scenario("bad", overrides, "x")
    assert message in str(exc.value.content)


def test_scenario_outside_benchmark_is_flagged_not_blocked(session):
    res = session.run_scenario("hot", {"entry_ev_multiple": 12.0, "exit_ev_multiple": 12.0}, "auction price")
    assert any("entry_ev_multiple" in n for n in res["benchmark_notes"])


def test_solve_for_target_hits_the_target(session):
    res = session.solve_for_target("entry_ev_multiple", "irr", 0.20, save_as="max_price_20")
    assert session.scenarios["max_price_20"].assumptions.entry_ev_multiple == pytest.approx(res["solution"], abs=1e-4)
    solved = session.scenarios["max_price_20"].assumptions
    assert run_model(solved)["returns"]["irr"] == pytest.approx(0.20, abs=1e-8)
    assert res["check"]["irr"] == pytest.approx(0.20, abs=1e-4)


def test_solve_for_exit_multiple_breakeven(session):
    res = session.solve_for_target("exit_ev_multiple", "moic", 1.0)
    a = session.apply_overrides(session.scenarios["base"].assumptions, {"exit_ev_multiple": res["solution"]})
    assert run_model(a)["returns"]["moic"] == pytest.approx(1.0, abs=1e-3)


def test_solve_for_unreachable_target_explains_range(session):
    with pytest.raises(ToolError) as exc:
        session.solve_for_target("senior_rate", "irr", 0.90)
    assert "not reachable" in str(exc.value.content)


def test_export_stays_in_output_dir_and_records_scenario(session, tmp_path):
    session.run_scenario("downside", {"revenue_growth": 0.0}, "flat sales in a recession")
    res = session.export_excel("downside", "../../etc/evil name.txt")
    path = res["path"]
    assert path == str(tmp_path / "out" / "evil_name.xlsx")
    wb = openpyxl.load_workbook(path)
    plan = {r[0].value: r for r in wb["Plan"].iter_rows(min_row=5) if r[0].value}
    assert plan["Revenue growth"][2].value.startswith("scenario (flat)")
    assert plan["Revenue growth"][3].value == 0.0                        # year 1 on the Plan sheet
    audit_text = [c.value for c in wb["Audit"]["C"] if isinstance(c.value, str)]
    assert "flat sales in a recession" in audit_text
    assert session.export_excel("base")["path"].endswith("Project_Test_base.xlsx")


def test_tools_before_base_case_ask_for_it(tmp_path):
    with pytest.raises(ToolError) as exc:
        DealSession(output_dir=str(tmp_path)).run_scenario("x", {}, "r")
    assert "generate_base_case" in exc.value.content


# --- agent loop, offline: scripted Claude responses through the real SDK ---------

def _message(content, stop_reason):
    return {"id": "msg", "type": "message", "role": "assistant", "model": "claude-opus-5", "content": content,
            "stop_reason": stop_reason, "stop_sequence": None, "usage": {"input_tokens": 10, "output_tokens": 10}}


def _tool_use(tid, name, args):
    return {"type": "tool_use", "id": tid, "name": name, "input": args}


class ScriptedClaude:
    """Replays a fixed sequence of API responses and records every request body."""

    def __init__(self, responses):
        self.responses = list(responses)
        self.requests = []

    def __call__(self, request):
        self.requests.append(json.loads(request.content))
        return httpx2.Response(200, json=self.responses.pop(0))

    def client(self):
        return anthropic.Anthropic(api_key="test", max_retries=0,
                                   http_client=DefaultHttpxClient(transport=httpx2.MockTransport(self)))


def _tool_results(request):
    return {b["tool_use_id"]: b for m in request["messages"] if m["role"] == "user" and isinstance(m["content"], list)
            for b in m["content"] if b.get("type") == "tool_result"}


def test_agent_loop_end_to_end(tmp_path):
    script = ScriptedClaude([
        _message([_tool_use("t1", "generate_base_case", {"description": "Valvole, fatturato 400m"})], "tool_use"),
        _message([  # parallel calls, one of them invalid
            _tool_use("t2", "run_scenario", {"name": "downside", "overrides": {"revenue_growth": 0.0},
                                             "rationale": "oil & gas capex cycle"}),
            _tool_use("t3", "run_scenario", {"name": "bad", "overrides": {"growth": 1}, "rationale": "typo"}),
            _tool_use("t4", "solve_for_target", {"variable": "entry_ev_multiple", "metric": "irr", "target": 0.2}),
        ], "tool_use"),
        _message([_tool_use("t5", "export_excel", {"scenario": "base"})], "tool_use"),
        _message([{"type": "text", "text": "Base case IRR ..."}], "end_turn"),
    ])
    calls = []
    session = DealSession(output_dir=str(tmp_path / "out"), llm=fake_llm)
    agent = LBOAgent(session, client=script.client(), on_tool_call=lambda n, a: calls.append(n))

    answer = agent.ask("Valvole, fatturato 400m. Prezzo massimo per IRR 20%?")

    assert answer == "Base case IRR ..."
    assert calls == ["generate_base_case", "run_scenario", "run_scenario", "solve_for_target", "export_excel"]
    assert set(session.scenarios) == {"base", "downside"}
    assert len(session.exports) == 1 and (tmp_path / "out").exists()

    first = script.requests[0]
    assert first["model"] == "claude-opus-5"
    assert first["system"] == SYSTEM_PROMPT
    assert [t["name"] for t in first["tools"]] == [
        "generate_base_case", "run_scenario", "solve_for_target", "compare_scenarios", "export_excel"]
    assert first["cache_control"] == {"type": "ephemeral"}
    assert first["fallbacks"] == "default"

    results = _tool_results(script.requests[2])
    assert "base_case_results" in results["t1"]["content"]
    assert not results["t2"].get("is_error")
    assert results["t3"]["is_error"] and "Unknown assumption" in results["t3"]["content"]
    assert json.loads(results["t4"]["content"])["solution"] > 0


def test_agent_follow_up_keeps_history(tmp_path):
    script = ScriptedClaude([
        _message([_tool_use("t1", "generate_base_case", {"description": "x"})], "tool_use"),
        _message([{"type": "text", "text": "first answer"}], "end_turn"),
        _message([_tool_use("t2", "run_scenario", {"name": "lev6", "overrides": {"total_leverage_x": 5.5},
                                                   "rationale": "user asked"})], "tool_use"),
        _message([{"type": "text", "text": "second answer"}], "end_turn"),
    ])
    agent = LBOAgent(DealSession(output_dir=str(tmp_path), llm=fake_llm), client=script.client())
    assert agent.ask("describe") == "first answer"
    assert agent.ask("and with 5.5x leverage?") == "second answer"
    follow_up_request = script.requests[2]
    roles = [m["role"] for m in follow_up_request["messages"]]
    assert roles == ["user", "assistant", "user", "assistant", "user"]
    assert "lev6" in agent.session.scenarios


def test_agent_surfaces_refusal(tmp_path):
    script = ScriptedClaude([{**_message([], "refusal"), "stop_details": {"type": "refusal", "category": None,
                                                                          "explanation": None}}])
    agent = LBOAgent(DealSession(output_dir=str(tmp_path), llm=fake_llm), client=script.client())
    with pytest.raises(RuntimeError, match="declined"):
        agent.ask("x")


def test_scenario_can_change_transaction_costs(session):
    res = session.run_scenario("pricier_debt", {"senior_oid_pct": 0.015, "fee_amortization_years": 5.0}, "tighter market")
    a = session.scenarios["pricier_debt"].assumptions
    assert a.senior_oid_pct == 0.015 and a.fee_amortization_years == 5
    base = session._summary("base", session.scenarios["base"].assumptions)
    assert res["entry_fees_and_oid"] > base["entry_fees_and_oid"] and res["irr"] < base["irr"]


def test_scenario_with_a_year_by_year_profile(session):
    res = session.run_scenario("gradual_erosion", {}, "margin erodes over three years",
                               plan={"ebitda_margin": [0.14, 0.12, 0.11]})
    a = session.scenarios["gradual_erosion"].assumptions
    assert a.ebitda_margin_by_year == (0.14, 0.12, 0.11)
    assert res["operating_plan"]["ebitda_margin"] == [0.14, 0.12, 0.11, 0.11, 0.11]
    assert res["overrides_vs_base"]["ebitda_margin_by_year"] == (0.14, 0.12, 0.11)
    assert res["entry_terms"].startswith("as signed")
    # a flat override afterwards replaces the profile
    session.run_scenario("flat_again", {"ebitda_margin": 0.13}, "flat", based_on="gradual_erosion")
    assert session.scenarios["flat_again"].assumptions.ebitda_margin_by_year is None


def test_plan_errors_are_actionable(session):
    with pytest.raises(ToolError) as exc:
        session.run_scenario("bad", {"ebitda_margin_by_year": 0.1}, "x")
    assert "`plan`" in exc.value.content
    with pytest.raises(ToolError) as exc:
        session.run_scenario("bad", {}, "x", plan={"nwc": [0.1]})
    assert "Unknown plan driver" in exc.value.content


def test_plan_scenario_exports_and_reloads(session, tmp_path):
    session.run_scenario("ramp", {}, "new plants ramp up", plan={"revenue_growth": [0.02, 0.06, 0.08]})
    path = session.export_excel("ramp")["path"]
    import openpyxl
    plan = {r[0].value: r for r in openpyxl.load_workbook(path)["Plan"].iter_rows(min_row=5) if r[0].value}
    assert [c.value for c in plan["Revenue growth"][3:8]] == [0.02, 0.06, 0.08, 0.08, 0.08]
    reloaded = DealSession.load(session.save(), str(tmp_path))
    assert reloaded.scenarios["ramp"].assumptions.revenue_growth_by_year == (0.02, 0.06, 0.08)


def test_tool_schema_exposes_plan():
    from lbo_agent import make_tools
    schema = next(t for t in make_tools(DealSession()) if t.to_dict()["name"] == "run_scenario").to_dict()
    assert "plan" in schema["input_schema"]["properties"]


def test_covenants_in_scenarios(session):
    base = session._summary("base", session.scenarios["base"].assumptions)
    assert base["covenants"]["first_breach_year"] is None                       # derived with headroom
    res = session.run_scenario("tight_bank", {"max_net_leverage": 3.0}, "lender asks for 3.0x")
    assert res["covenants"]["first_breach_year"] == 1
    assert res["covenants"]["max_net_leverage_by_year"][0] == 3.0             # flat value replaced the steps
    assert any(w.startswith("Covenant breach in year 1") for w in res["engine_warnings"])
    res = session.run_scenario("steps", {}, "custom step-downs", plan={"max_net_leverage": [6.0, 5.5, 5.0]})
    assert res["covenants"]["max_net_leverage_by_year"][:3] == [6.0, 5.5, 5.0]


def test_circularity_switch_in_scenarios(session):
    res = session.run_scenario("average_interest", {"interest_on_average_balance": 1}, "desk convention")
    a = session.scenarios["average_interest"].assumptions
    assert a.interest_on_average_balance is True
    base = session._summary("base", session.scenarios["base"].assumptions)
    assert res["irr"] != base["irr"]
    assert res["overrides_vs_base"] == {"interest_on_average_balance": True}
