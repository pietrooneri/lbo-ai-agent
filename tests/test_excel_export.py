"""The exported workbook is a second implementation of the engine, in Excel formulas.
These tests evaluate it (pycel) and require it to match lbo_engine.py line by line."""

import dataclasses
import math
import re

import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from pycel import ExcelCompiler

from assumption_generator import apply_guardrails
from excel_export import ASSUMPTION_LAYOUT, PLAN_FIRST_ROW, export_to_excel, lbo_rows
from lbo_engine import PLAN_DRIVERS
from lbo_engine import Assumptions, run_model
from test_assumption_generator import proposal
from test_lbo_engine import SCENARIOS

YEAR_LINES = {  # LBO row key -> YearResult attribute
    "revenue": "revenue", "ebitda": "ebitda", "da": "da", "ebit": "ebit", "interest": "interest_expense",
    "ebt": "ebt", "tax": "tax", "ni": "net_income", "capex": "capex", "nwc": "nwc_change",
    "fcf": "fcf_pre_sweep", "mand": "mandatory_amort", "sweep": "cash_sweep",
    "rcf_draw": "rcf_draw", "rcf_repay": "rcf_repayment", "tl_end": "senior_end_balance",
    "sub_end": "sub_end_balance", "rcf_end": "rcf_end_balance", "cash_end": "cash_end_balance",
}


def _assumption_cell(name):
    row = 5
    for field, label, *_ in ASSUMPTION_LAYOUT:
        if label is None:
            row += 2
            continue
        if field == name:
            return f"Assumptions!B{row}"
        row += 1
    raise KeyError(name)


class Book:
    def __init__(self, path, a):
        self.xl = ExcelCompiler(filename=str(path), plugins=("pycel_plugins",))
        self.R = lbo_rows(a)

    def evaluate(self, address):
        return self.xl.evaluate(address)

    def lbo(self, key, col="C"):
        return self.evaluate(f"LBO!{col}{self.R[key]}")

    def year(self, key, t):
        return self.lbo(key, get_column_letter(3 + t))

    def set(self, name, value):
        self.xl.evaluate("Audit!C10")  # pycel only lets you edit cells it has already loaded
        if name in PLAN_DRIVERS:       # year-by-year driver: set every year on the Plan sheet
            row = PLAN_FIRST_ROW + PLAN_DRIVERS.index(name)
            for t in range(1, 11):
                self.xl.set_value(f"Plan!{get_column_letter(3 + t)}{row}", value)
        else:
            self.xl.set_value(_assumption_cell(name), value)


def _export(tmp_path, a, audit=None, circular=False):
    """circular=False: the loop-free twin of the workbook (opening-balance interest formulas), which
    pycel can evaluate; everything else is identical to the file users get. The circular interest
    formulas are covered by the tests at the end of this file and by a check in Microsoft Excel."""
    path = tmp_path / "lbo.xlsx"
    export_to_excel(a, str(path), audit, circular=circular)
    return Book(path, a)


def _assert_matches_engine(book, a):
    out = run_model(a)
    for t, yr in enumerate(out["years"], start=1):
        for key, attr in YEAR_LINES.items():
            assert book.year(key, t) == pytest.approx(getattr(yr, attr), abs=1e-6), f"{key} year {t}"
    ret = out["returns"]
    assert book.lbo("su_equity") == pytest.approx(out["sources_uses"]["sources"]["sponsor_equity"])
    assert book.lbo("exit_equity") == pytest.approx(ret["exit_equity_value"], abs=1e-6)
    assert book.lbo("moic") == pytest.approx(ret["moic"], abs=1e-9)
    assert book.lbo("irr") == pytest.approx(ret["irr"], abs=1e-9)
    assert book.lbo("irr_chk") == pytest.approx(ret["irr"], abs=1e-7)
    assert book.lbo("peak_rcf") == pytest.approx(ret["peak_rcf_draw"], abs=1e-6)
    if math.isfinite(ret["min_interest_coverage"]):
        assert book.lbo("min_cov") == pytest.approx(ret["min_interest_coverage"], abs=1e-9)
    assert (book.lbo("flag_rcf") != "-") == any("RCF" in w for w in ret["warnings"])
    assert (book.lbo("flag_cov") != "-") == any("coverage" in w for w in ret["warnings"])
    assert book.evaluate("LBO!C3") == "OK"
    cov = ret["covenants"]
    if cov:
        for t in cov["tests"]:
            for key, row in (("leverage_headroom", "cv_lev_head"), ("cover_headroom", "cv_cov_head")):
                if t[key] is not None:
                    assert book.year(row, t["year"]) == pytest.approx(t[key], abs=1e-9), (key, t["year"])
            assert book.year("cv_test", t["year"]) == ("BREACH" if t["breach"] else "OK")
        assert book.lbo("cv_first") == (cov["first_breach_year"] or "none")
        assert book.lbo("cv_min_head") == pytest.approx(cov["min_headroom"], abs=1e-9)
    assert (book.lbo("flag_cov_breach") != "-") == any("Covenant breach" in w for w in ret["warnings"])


OPENING_BALANCE_SCENARIOS = [n for n, a in SCENARIOS.items() if not a.interest_on_average_balance]


@pytest.mark.parametrize("name", OPENING_BALANCE_SCENARIOS)
def test_workbook_matches_engine(tmp_path, name):
    a = SCENARIOS[name]
    _assert_matches_engine(_export(tmp_path, a), a)


def test_workbook_from_generator_output(tmp_path):
    res = apply_guardrails(proposal(provided={"revenue_at_entry"}, ebitda_margin=0.35))
    book = _export(tmp_path, res.assumptions, res.to_dict())
    _assert_matches_engine(book, res.assumptions)
    assert book.evaluate("Audit!C10").startswith("Matches the Python engine")
    wb = openpyxl.load_workbook(tmp_path / "lbo.xlsx")
    notes = [c.value for c in wb["Plan"]["C"] if c.value]           # margin now lives on the Plan sheet
    assert any("Guardrail" in n and "outside industrials margin range" in n for n in notes)


@pytest.mark.parametrize("name, value", [
    ("hold_period_years", 7), ("revenue_growth", 0.08), ("ebitda_margin", 0.12), ("entry_ebitda", 140.0),
    ("total_leverage_x", 6.0), ("exit_ev_multiple", 9.5), ("cash_sweep_pct", 0.5),
])
def test_editing_an_input_in_excel_reprices_the_model(tmp_path, name, value):
    """The point of 'live formulas': change one blue cell, get the engine's answer for it."""
    base = SCENARIOS["base"]
    book = _export(tmp_path, base)
    book.set(name, value)
    flexed = dataclasses.replace(base, **{name: value})
    ret = run_model(flexed)["returns"]
    assert book.lbo("moic") == pytest.approx(ret["moic"], abs=1e-9)
    assert book.lbo("irr") == pytest.approx(ret["irr"], abs=1e-9)
    assert book.evaluate("Audit!C10").startswith("Inputs changed")


def test_sensitivity_table_matches_engine(tmp_path):
    base = SCENARIOS["base"]
    book = _export(tmp_path, base)
    R = book.R
    sens_hdr = max(R.values()) + 3            # section title, then header row
    for i, off in enumerate([-1.0, -0.5, 0.0, 0.5, 1.0]):
        for j, hold in enumerate([3, 4, 5, 6, 7]):
            cell = f"LBO!{get_column_letter(4 + j)}{sens_hdr + 1 + i}"
            flexed = dataclasses.replace(base, exit_ev_multiple=base.exit_ev_multiple + off, hold_period_years=hold)
            assert book.evaluate(cell) == pytest.approx(run_model(flexed)["returns"]["irr"], abs=1e-9), cell


def test_lbo_sheet_has_no_hardcoded_numbers(tmp_path):
    """Only structural zeros (entry year index, opening RCF) and the sensitivity axis are constants."""
    a = SCENARIOS["base"]
    export_to_excel(a, str(tmp_path / "lbo.xlsx"))
    ws = openpyxl.load_workbook(tmp_path / "lbo.xlsx")["LBO"]
    R = lbo_rows(a)
    sens_rows = range(max(R.values()) + 3, ws.max_row + 1)
    constants = [c.coordinate for row in ws.iter_rows() for c in row
                 if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
                 and c.row not in sens_rows]
    assert sorted(constants) == sorted([f"C{R['years']}", f"C{R['rcf_end']}"])


def test_editing_one_year_of_the_plan_reprices_that_year_only(tmp_path):
    base = SCENARIOS["base"]
    book = _export(tmp_path, base)
    book.evaluate("Audit!C10")
    book.xl.set_value(f"Plan!F{PLAN_FIRST_ROW + 1}", 0.15)          # margin, year 3 only
    flexed = dataclasses.replace(base, ebitda_margin_by_year=[0.20, 0.20, 0.15, 0.20])
    assert book.lbo("moic") == pytest.approx(run_model(flexed)["returns"]["moic"], abs=1e-9)


@pytest.mark.parametrize("name", ["average_interest", "average_interest_rcf"])
def test_average_balance_interest_formulas(tmp_path, name):
    """Circularity switch on. pycel does not solve the whole reference loop consistently (its
    total can lag its components), so this checks each interest formula at pycel's state:
    interest = rate x (opening + closing) / 2 for every tranche and year. The full numbers are
    checked against the engine in Microsoft Excel (see README)."""
    a = SCENARIOS[name]
    book = _export(tmp_path, a, circular=True)
    book.lbo("moic")
    book.xl.recalculate()
    for t in range(1, a.hold_period_years + 1):
        for line, balance, rate in (("int_senior", "tl_end", a.senior_rate), ("int_sub", "sub_end", a.sub_rate),
                                    ("int_rcf", "rcf_end", a.rcf_rate)):
            opening, closing = book.year(balance, t - 1) if t > 1 else book.lbo(balance), book.year(balance, t)
            assert book.year(line, t) == pytest.approx(rate * (opening + closing) / 2, abs=1e-6), (line, t)


def test_workbook_has_iterative_calculation_and_switch(tmp_path):
    export_to_excel(SCENARIOS["average_interest"], str(tmp_path / "lbo.xlsx"))
    wb = openpyxl.load_workbook(tmp_path / "lbo.xlsx")
    assert wb.calculation.iterate and wb.calculation.iterateCount >= 100
    assert wb["Assumptions"][_assumption_cell("interest_on_average_balance").split("!")[1]].value == 1


def test_user_workbook_interest_formulas_carry_the_switch(tmp_path):
    """The file users get: each interest line averages opening and closing balance when the switch
    is 1 and uses the opening balance when it is 0; the loop-free twin differs only there."""
    a = SCENARIOS["base"]
    export_to_excel(a, str(tmp_path / "user.xlsx"))
    export_to_excel(a, str(tmp_path / "twin.xlsx"), circular=False)
    user, twin = (openpyxl.load_workbook(tmp_path / f)["LBO"] for f in ("user.xlsx", "twin.xlsx"))
    R, switch = lbo_rows(a), _assumption_cell("interest_on_average_balance")
    switch = switch.replace("Assumptions!B", "Assumptions!$B$")
    for line, bal in (("int_senior", "tl_end"), ("int_sub", "sub_end"), ("int_rcf", "rcf_end")):
        f = user[f"E{R[line]}"].value
        assert f.startswith(f"=IFERROR(IF({switch}=1,(D{R[bal]}+E{R[bal]})/2,D{R[bal]}),D{R[bal]})*"), f
        assert twin[f"E{R[line]}"].value == "=D{}*{}".format(R[bal], f.split(")*")[1])
    differing = [c.coordinate for row in user.iter_rows() for c in row
                 if c.value != twin[c.coordinate].value]
    assert {int(re.sub("[A-Z]", "", x)) for x in differing} == {R["int_senior"], R["int_sub"], R["int_rcf"]}
